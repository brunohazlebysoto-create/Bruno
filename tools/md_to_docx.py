"""Converts structured medical Markdown to a formatted Word (.docx) document."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY  = RGBColor(0x0A, 0x29, 0x5C)
BLUE  = RGBColor(0x1A, 0x6B, 0xC4)
CYAN  = RGBColor(0x00, 0xB4, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0x44, 0x44, 0x55)

_BOLD_RE  = re.compile(r'\*\*(.+?)\*\*')
_TABLE_RE = re.compile(r'^\|')
_SEP_RE   = re.compile(r'^\|[-| :]+\|$')


def _set_cell_bg(cell, color: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    hex_ = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_)
    tcPr.append(shd)


def _add_styled_run(para, text: str, bold=False, size_pt: int = 11,
                    color: RGBColor = None):
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    return run


def _add_inline(para, line: str, base_size: int = 11, base_color=None):
    """Add a line with inline **bold** support."""
    parts = _BOLD_RE.split(line)
    for i, part in enumerate(parts):
        if not part:
            continue
        is_bold = (i % 2 == 1)
        _add_styled_run(para, part, bold=is_bold,
                        size_pt=base_size,
                        color=base_color if not is_bold else None)


def _style_document(doc: Document):
    """Apply global document styles."""
    style = doc.styles['Normal']
    style.font.name  = 'Calibri'
    style.font.size  = Pt(11)
    style.font.color.rgb = GRAY

    for lvl, size, color, space_before in [
        ('Heading 1', 18, NAVY, 18),
        ('Heading 2', 14, BLUE, 14),
        ('Heading 3', 12, BLUE, 10),
        ('Heading 4', 11, GRAY, 8),
    ]:
        s = doc.styles[lvl]
        s.font.name  = 'Calibri'
        s.font.size  = Pt(size)
        s.font.color.rgb = color
        s.font.bold  = True
        s.paragraph_format.space_before = Pt(space_before)
        s.paragraph_format.space_after  = Pt(4)


def _collect_table(lines: list, start: int):
    """Collect consecutive table lines starting at start."""
    rows = []
    i = start
    while i < len(lines) and _TABLE_RE.match(lines[i].strip()):
        raw = lines[i].strip()
        if not _SEP_RE.match(raw):
            cols = [c.strip() for c in raw.strip('|').split('|')]
            rows.append(cols)
        i += 1
    return rows, i


def _add_table(doc: Document, rows: list):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    # Pad rows
    rows = [r + [''] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    for ri, row_data in enumerate(rows):
        row = table.rows[ri]
        for ci, cell_text in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = ''
            para = cell.paragraphs[0]
            if ri == 0:
                _set_cell_bg(cell, NAVY)
                run = para.add_run(cell_text)
                run.bold = True
                run.font.color.rgb = WHITE
                run.font.size = Pt(10)
            else:
                _add_inline(para, cell_text, base_size=10)


def md_to_docx(md_text: str, output_path: str):
    """Convert a Markdown string to a styled Word document."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _style_document(doc)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ── Headings ───────────────────────────────────────────────────────── #
        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
            i += 1; continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            i += 1; continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            i += 1; continue
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
            i += 1; continue

        # ── Table ─────────────────────────────────────────────────────────── #
        if _TABLE_RE.match(stripped):
            rows, i = _collect_table(lines, i)
            _add_table(doc, rows)
            doc.add_paragraph()
            continue

        # ── Bullet list ───────────────────────────────────────────────────── #
        if re.match(r'^[-*] ', stripped):
            para = doc.add_paragraph(style='List Bullet')
            _add_inline(para, stripped[2:])
            i += 1; continue

        # ── Numbered list ─────────────────────────────────────────────────── #
        if re.match(r'^\d+\. ', stripped):
            para = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\. ', '', stripped)
            _add_inline(para, text)
            i += 1; continue

        # ── Horizontal rule ───────────────────────────────────────────────── #
        if stripped in ('---', '***', '___'):
            doc.add_paragraph('─' * 60)
            i += 1; continue

        # ── Empty line ────────────────────────────────────────────────────── #
        if not stripped:
            i += 1; continue

        # ── Normal paragraph ──────────────────────────────────────────────── #
        para = doc.add_paragraph()
        _add_inline(para, stripped)
        i += 1

    doc.save(output_path)
